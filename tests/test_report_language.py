"""
tests/test_report_language.py

Guards against authority-voice language ("The Tribunal hereby declares...",
"the Tribunal awards...", etc.) leaking into generated reports. This is the
Sprint 1 fix for the lawyer's core feedback: the system must never frame its
output as an actual adjudicator's ruling.

Two layers:
1. Direct check against the fallback award_framework templates (fast,
   deterministic, no LLM calls — mocks _call_gemini to force the except path).
2. Full .docx generation + read-back for scenario_1 (NOT ARBITRABLE) and
   scenario_2 (ARBITRABLE), to catch banned phrases anywhere in the actual
   output file, not just the template strings.
"""

import glob
import json
import os
import re
import sys
from unittest.mock import patch

import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from agents.arbitrability_agent import check_arbitrability
from agents.gemini_agents import generate_award_framework
from agents.landmark_retrieval_agent import retrieve_landmarks, analyze_landmark_applicability
from agents.report_generator import generate_dss_report

# ---------------------------------------------------------------------------
# Banned phrase patterns — authority-voice language that must never appear.
# Case-insensitive; keep this list in sync with whatever the Sprint 1 rewrite
# prompt targets in gemini_agents.py.
# ---------------------------------------------------------------------------
BANNED_PATTERNS = [
    r"\bthe\s+tribunal\s+hereby\s+declares\b",
    r"\bthe\s+tribunal\s+finds\b",
    r"\bthe\s+tribunal\s+awards\b",
    r"\bthe\s+tribunal\s+holds\b",
    r"\bthe\s+tribunal\s+\[holds/does\s+not\s+hold\]\b",
    r"\bin\s+the\s+matter\s+of\s+arbitration\s+between\b",
    r"\bsole\s+arbitrator\b",
]
BANNED_RE = re.compile("|".join(BANNED_PATTERNS), flags=re.IGNORECASE)


def _find_banned(text: str) -> list[str]:
    """Return every banned phrase match found in text (empty list = clean)."""
    if not text:
        return []
    return [m.group(0) for m in BANNED_RE.finditer(text)]


def _flatten_award_framework_text(framework: dict) -> str:
    """Concatenate every string field in an award_framework dict for scanning."""
    parts = [framework.get("jurisdiction_finding", "")]
    for finding in framework.get("findings_on_issues", []):
        parts.append(finding.get("issue", ""))
        parts.extend(finding.get("finding_options", []))
        parts.append(finding.get("applicable_law", ""))
    relief = framework.get("relief_section", {})
    parts.append(relief.get("injunction_guidance", ""))
    parts.append(relief.get("damages_guidance", ""))
    parts.append(relief.get("costs_guidance", ""))
    parts.append(framework.get("operative_portion_template", ""))
    return "\n".join(p for p in parts if p)


# ---------------------------------------------------------------------------
# Layer 1 — fallback template check (no LLM calls, forces the except path)
# ---------------------------------------------------------------------------

_SCENARIO_DIR = os.path.join(os.path.dirname(__file__), "test_cases")


def _load_scenario(filename: str) -> dict:
    with open(os.path.join(_SCENARIO_DIR, filename), encoding="utf-8") as f:
        return json.load(f)


@pytest.mark.parametrize(
    "scenario_file",
    ["scenario_1.json", "scenario_2.json"],  # NOT ARBITRABLE, ARBITRABLE
)
def test_fallback_award_framework_has_no_authority_voice(scenario_file):
    """With Gemini/Groq/LM Studio all forced to fail, the hardcoded fallback
    template in generate_award_framework()'s except block must still be
    free of authority-voice language."""
    dispute = _load_scenario(scenario_file)
    arbitrability_result = check_arbitrability(dispute)
    landmark_matches = retrieve_landmarks(
        dispute["dispute_description"],
        arbitrability_result,
        dispute_type_label=dispute.get("dispute_type", ""),
    )

    with patch("agents.gemini_agents._call_gemini", side_effect=Exception("forced failure")):
        framework = generate_award_framework(
            dispute,
            extracted_facts={},
            arbitrability_result=arbitrability_result,
            issues=[
                "Whether the mark is deceptively similar?",
                "Whether infringement occurred?",
                "Whether injunctive relief is warranted?",
                "Whether damages are payable?",
            ],
            principles=[],
        )

    assert framework.get("generation_method") == "fallback", (
        "Test setup error: expected the fallback path to fire."
    )

    full_text = _flatten_award_framework_text(framework)
    hits = _find_banned(full_text)
    assert not hits, (
        f"Authority-voice language found in fallback award_framework for "
        f"{scenario_file}: {hits}\n\nFull text:\n{full_text}"
    )


# ---------------------------------------------------------------------------
# Layer 2 — full .docx generation + read-back
# ---------------------------------------------------------------------------

def _extract_all_text(docx_path: str) -> str:
    """Pull every paragraph and table-cell string out of a generated .docx."""
    doc = Document(docx_path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


@pytest.mark.parametrize(
    "scenario_file",
    ["scenario_1.json", "scenario_2.json"],
)
def test_generated_docx_has_no_authority_voice(scenario_file, tmp_path, monkeypatch):
    """End-to-end: generate an actual report (LLM calls forced to fail, so
    this stays free and deterministic) and scan the resulting .docx file
    for banned phrases anywhere in the document, including tables."""
    dispute = _load_scenario(scenario_file)
    arbitrability_result = check_arbitrability(dispute)
    landmark_matches = retrieve_landmarks(
        dispute["dispute_description"],
        arbitrability_result,
        dispute_type_label=dispute.get("dispute_type", ""),
    )
    landmark_analyses = [
        analyze_landmark_applicability(dispute, lm) for lm in landmark_matches
    ]

    with patch("agents.gemini_agents._call_gemini", side_effect=Exception("forced failure")):
        framework = generate_award_framework(
            dispute,
            extracted_facts={},
            arbitrability_result=arbitrability_result,
            issues=[
                "Whether the mark is deceptively similar?",
                "Whether infringement occurred?",
                "Whether injunctive relief is warranted?",
                "Whether damages are payable?",
            ],
            principles=[],
        )

    # Redirect OUTPUT_DIR to a pytest tmp_path so this test doesn't litter
    # the real output/ folder.
    import agents.report_generator as rg
    monkeypatch.setattr(rg, "OUTPUT_DIR", str(tmp_path))

    filepath = generate_dss_report(
        dispute=dispute,
        extracted_facts={},
        arbitrability_result=arbitrability_result,
        landmark_matches=landmark_matches,
        landmark_analyses=landmark_analyses,
        issues=[
            "Whether the mark is deceptively similar?",
            "Whether infringement occurred?",
            "Whether injunctive relief is warranted?",
            "Whether damages are payable?",
        ],
        legal_principles=[],
        award_framework=framework,
        adversarial_analysis=None,
        generation_methods={"framework": "fallback"},
    )

    full_text = _extract_all_text(filepath)
    hits = _find_banned(full_text)
    assert not hits, (
        f"Authority-voice language found in generated .docx for "
        f"{scenario_file}: {hits}"
    )


# ---------------------------------------------------------------------------
# Standalone sanity check — run directly for a quick manual check without pytest
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    print("Run via pytest: python -m pytest tests/test_report_language.py -v")